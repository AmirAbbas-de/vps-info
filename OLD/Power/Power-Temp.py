import hpilo
from ipaddress import IPv4Network, IPv4Address
from concurrent.futures import ThreadPoolExecutor
from ping3 import ping
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import json
import sys
import os

def load_config(dc_name):
    filename = os.path.join(os.path.dirname(__file__), f"{dc_name}.json")
    if not os.path.exists(filename):
        raise FileNotFoundError(f"Config file {filename} not found.")
    with open(filename, "r") as f:
        data = json.load(f)
    return data["ranges"]

wb = openpyxl.Workbook()
summary_ws = wb.active
summary_ws.title = "Summary"

headers = [
    "iLO IP",
    "Ambient (°C)", "CPU 1 (°C)", "CPU 2 (°C)",
    "Current Power (W)", "Min Power (W)", "Avg Power (W)", "Max Power (W)",
    "Status"
]

green_fill = PatternFill(start_color='00FF00', end_color='00FF00', fill_type='solid')
yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
red_fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
light_blue_fill = PatternFill(start_color='6df7c7', end_color='6df7c7', fill_type='solid')
navi_blue_fill = PatternFill(start_color='00b3b3', end_color='00b3b3', fill_type='solid')

rack_results = {}
summary_results = []


def check_ping(ip):
    """Check if IP responds to ping with multiple attempts"""
    try:
        attempts = 3
        timeout = 0.1
        ip_str = str(ip)
        failed_attempts = 0
        
        for _ in range(attempts):
            response = ping(ip_str, timeout=timeout)
            if response is None or response is False:
                failed_attempts += 1
        
        if failed_attempts == attempts:
            return None
        return ip_str
    except Exception as e:
        print(f"❌ Error pinging {ip}: {e}")
        return None


def process_ip(ip):
    ip_str = str(ip)
    third_octet = ip_str.split(".")[2]
    rack_key = f"Rack-{third_octet}"
    status = "OK"
    ambient = "N/A"
    cpu1 = "N/A"
    cpu2 = "N/A"
    curr_power = "N/A"
    min_power = "N/A"
    avg_power = "N/A"
    max_power = "N/A"

    try:
        iloUsername = "s.miran"
        iLOpass = "5593Posa@@"
        ilo = hpilo.Ilo(ip_str, iloUsername, iLOpass)
    except Exception as e:
        status = f"Login failed: {str(e)}"
        print(f"❌ Error logging into iLO {ip_str}: {e}")
    else:
        try:
            data = ilo.get_embedded_health()
            temperatures = data.get('temperature', {})

            required_sensors = ['01-Inlet Ambient', '02-CPU 1', '03-CPU 2']
            for i, sensor in enumerate(required_sensors):
                values = temperatures.get(sensor, {})
                current_reading = values.get('currentreading', ('N/A', ''))
                temp = current_reading[0] if isinstance(current_reading, tuple) else 'N/A'
                if i == 0:
                    ambient = temp
                elif i == 1:
                    cpu1 = temp
                elif i == 2:
                    cpu2 = temp
        except Exception as e:
            error_detail = f"Temp failed: {str(e)}"
            status = status + " | " + error_detail if status != "OK" else error_detail
            print(f"❌ Error fetching temp data for {ip_str}: {e}")

        try:
            power_data = ilo.get_power_readings()
            curr_power = power_data.get('present_power_reading', ["N/A"])[0]
            min_power = power_data.get('minimum_power_reading', ["N/A"])[0]
            avg_power = power_data.get('average_power_reading', ["N/A"])[0]
            max_power = power_data.get('maximum_power_reading', ["N/A"])[0]
        except Exception as e:
            error_detail = f"Power failed: {str(e)}"
            status = status + " | " + error_detail if status != "OK" else error_detail
            print(f"❌ Error fetching power data for {ip_str}: {e}")

    row = [ip_str, ambient, cpu1, cpu2, curr_power, min_power, avg_power, max_power, status]
    rack_results.setdefault(rack_key, []).append(row)
    summary_results.append(row)


def colorize_row(sheet, row_number):
    try:
        rules = {
            2: [(None, 30, green_fill), (30, 35, yellow_fill), (35, None, red_fill)],  # Ambient
            3: [(None, 40, green_fill), (40, 50, yellow_fill), (50, None, red_fill)],  # CPU 1
            4: [(None, 40, green_fill), (40, 50, yellow_fill), (50, None, red_fill)],  # CPU 2
            5: [(None, 150, green_fill), (150, 219, yellow_fill), (219, None, red_fill)],  # Current Power
            6: [(None, 150, green_fill), (150, 219, yellow_fill), (219, None, red_fill)],  # Min Power
            7: [(None, 150, green_fill), (150, 219, yellow_fill), (219, None, red_fill)],  # Avg Power
            8: [(None, 150, green_fill), (150, 219, yellow_fill), (219, None, red_fill)],  # Max Power
        }

        for col_index, ranges in rules.items():
            cell = sheet.cell(row=row_number, column=col_index)
            value = cell.value
            if isinstance(value, (int, float)):
                for lower, upper, fill in ranges:
                    if (lower is None or value >= lower) and (upper is None or value < upper):
                        cell.fill = fill
                        break
    except Exception as e:
        print(f"⚠️ Error coloring row {row_number}: {e}")


def format_worksheet(ws, rows, sheet_name):
    sorted_rows = sorted(rows, key=lambda x: IPv4Address(x[0]))
    ws.append(headers)
    for row in sorted_rows:
        ws.append(row)
        colorize_row(ws, ws.max_row)
    for col in ws.columns:
        max_length = max((len(str(cell.value)) for cell in col if cell.value), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_length + 2
    table_range = f"A1:I{ws.max_row}"
    table_name = f"Table_{sheet_name.replace('-', '_')[:25]}"
    tab = Table(displayName=table_name, ref=table_range)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    ws.add_table(tab)


def add_power_sums(ws):
    total_row = ws.max_row + 1
    normalized_row = total_row + 1

    total_cell = ws.cell(row=total_row, column=1, value="Total by Watts")
    total_cell.font = Font(bold=True)

    normalized_cell = ws.cell(row=normalized_row, column=1, value="Total by Amps")
    normalized_cell.font = Font(italic=True)

    for col in range(5, 9):
        col_letter = get_column_letter(col)
        data_range = f"{col_letter}2:{col_letter}{total_row - 1}"

        total_cell = ws.cell(row=total_row, column=col, value=f"=SUM({data_range})")
        total_cell.fill = light_blue_fill

        normalized_cell = ws.cell(row=normalized_row, column=col, value=f"={col_letter}{total_row}/220")
        normalized_cell.number_format = '0.00'
        normalized_cell.fill = navi_blue_fill


def ilo(dc_name):
    ranges = load_config(dc_name)
    all_ips = [str(ip) for r in ranges for ip in IPv4Network(r).hosts() if str(ip).split('.')[-1] != '1']

    with ThreadPoolExecutor(max_workers=60) as executor:
        ping_results = list(executor.map(check_ping, all_ips))
    
    responsive_ips = [ip for ip in ping_results if ip is not None]
    
    with ThreadPoolExecutor(max_workers=60) as executor:
        executor.map(process_ip, responsive_ips)

    format_worksheet(summary_ws, summary_results, "Summary")
    add_power_sums(summary_ws)

    for rack, rows in rack_results.items():
        ws = wb.create_sheet(title=rack)
        format_worksheet(ws, rows, rack)
        add_power_sums(ws)

    rack_sheets = [sheet for sheet in wb.sheetnames if sheet.startswith("Rack-")]
    sorted_racks = sorted(rack_sheets, key=lambda name: int(name.split("-")[1]))
    wb._sheets = [wb["Summary"]] + [wb[sheet] for sheet in sorted_racks]

    file_name = f"{dc_name}-Temprature_PowerUsage.xlsx"
    wb.save(file_name)
    print(f"\nTemp + Power report (Summary + Racks) saved to {file_name}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("❌ Please provide the datacenter name")
        sys.exit(1)

    dc_name = sys.argv[1]
    ilo(dc_name)

